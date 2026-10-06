"""Tests that need no API key, no database and no network.

    uv run python -m unittest discover -s tests -v
"""

import json
import tempfile
import unittest
from pathlib import Path

from pydantic import ValidationError

from graph_nodes.export import slugify, write_code_file, write_quiz_files
from MCP_Servers.pptx_builder import build_presentation
from schemas import ApproveRequest, QuizQuestion, SessionRequest
from Session_State import next_agent

BRIEF = {"lecturer_id": "Dr. Sara", "topic": "Python decorators", "needs": ["slides", "quiz"]}
QUIZ = [{"question": "What does @ do?", "options": ["a", "b", "c", "d"],
         "answer_index": 1, "explanation": "It applies a decorator."}]


class RoutingTests(unittest.TestCase):
    def test_skips_agents_that_were_not_selected(self):
        state = {"selected_agents": ["research_agent", "quiz_agent"]}
        self.assertEqual(next_agent(state, "orchestrator"), "research_agent")
        self.assertEqual(next_agent(state, "research_agent"), "quiz_agent")
        self.assertEqual(next_agent(state, "quiz_agent"), "draft_agent")

    def test_nothing_selected_goes_to_draft(self):
        self.assertEqual(next_agent({"selected_agents": []}, "orchestrator"), "draft_agent")


class RequestModelTests(unittest.TestCase):
    def test_objectives_text_becomes_a_list(self):
        req = SessionRequest(**BRIEF, learning_objectives="- write a decorator\n\nexplain closures ")
        self.assertEqual(req.learning_objectives, ["write a decorator", "explain closures"])

    def test_empty_objectives(self):
        self.assertEqual(SessionRequest(**BRIEF, learning_objectives="").learning_objectives, [])

    def test_rejects_unknown_need_and_empty_needs(self):
        with self.assertRaises(ValidationError):
            SessionRequest(**{**BRIEF, "needs": ["video"]})
        with self.assertRaises(ValidationError):
            SessionRequest(**{**BRIEF, "needs": []})

    def test_revise_requires_feedback(self):
        with self.assertRaises(ValidationError):
            ApproveRequest(session_id="abc", decision="revise", feedback="  ")
        ApproveRequest(session_id="abc", decision="approve")

    def test_quiz_question_must_have_four_options(self):
        with self.assertRaises(ValidationError):
            QuizQuestion(question="q", options=["a", "b"], answer_index=0, explanation="e")
        with self.assertRaises(ValidationError):
            QuizQuestion(question="q", options=["a", "b", "c", "d"], answer_index=4, explanation="e")


class ExportTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_slugify_is_safe_for_file_names(self):
        self.assertEqual(slugify("Python Decorators!"), "python-decorators")
        self.assertEqual(slugify("../../etc/passwd"), "etc-passwd")
        self.assertEqual(slugify("شرح", "slides"), "slides")  # no Latin letters: use the fallback

    def test_code_file_is_valid_python_and_reports_status(self):
        examples = [
            {"title": "Hello", "code": "print('hi')", "explanation": "Prints hi.", "verified": True},
            {"title": "Broken", "code": "print(1)", "explanation": "x", "verified": False},
        ]
        path = write_code_file(self.folder, examples, "Python")
        text = path.read_text(encoding="utf-8")
        self.assertEqual(path.name, "examples.py")
        compile(text, str(path), "exec")  # the comments must not break the file
        self.assertIn("verified by running it", text)
        self.assertIn("NOT verified", text)

    def test_code_file_extension_for_other_languages(self):
        ex = [{"title": "t", "code": "int x;", "explanation": "e", "verified": None}]
        self.assertEqual(write_code_file(self.folder, ex, "C++").name, "examples.cpp")

    def test_quiz_files(self):
        json_path, md_path = write_quiz_files(self.folder, QUIZ, "Decorators")
        self.assertEqual(json.loads(json_path.read_text(encoding="utf-8"))["questions"], QUIZ)
        md = md_path.read_text(encoding="utf-8")
        self.assertIn("B. b", md)
        self.assertIn("1. B: It applies a decorator.", md)

    @staticmethod
    def _texts(slide) -> list[str]:
        return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text]

    def test_pptx_structure_footer_and_notes(self):
        from pptx import Presentation

        deck = {
            "title": "Decorators", "course_name": "Advanced Python", "lecturer": "Dr. Sara",
            "slides": [
                {"title": "Closures", "bullets": ["one", "two"], "notes": "Say this.",
                 "diagram_steps": ["Define", "Wrap", "Call"]},
                {"title": "Summary", "bullets": ["done"], "notes": ""},
            ],
            "code_examples": [{"title": "Hello", "code": "print('hi')", "explanation": "Prints.", "output": "hi"}],
            "quiz": QUIZ,
        }
        prs = Presentation(str(build_presentation(self.folder / "deck.pptx", deck)))
        slides = list(prs.slides)
        # cover, Closures, code, question, answer, Summary
        self.assertEqual(len(slides), 6)
        self.assertIn("Decorators", self._texts(slides[0]))
        closures = self._texts(slides[1])
        for expected in ["Closures", "one", "two", "Define", "Wrap", "Call"]:
            self.assertIn(expected, closures)
        self.assertEqual(slides[1].notes_slide.notes_text_frame.text, "Say this.")
        self.assertTrue(any("print('hi')" in t for t in self._texts(slides[2])))
        self.assertIn("Summary", self._texts(slides[5]))  # the summary closes the deck
        # the lecturer's name and the slide number are on every slide after the cover
        for number, slide in enumerate(slides[1:], 2):
            texts = self._texts(slide)
            self.assertIn("Dr. Sara · Advanced Python", texts)
            self.assertIn(f"{number} / 6", texts)

    def test_pptx_with_slides_only_and_arabic_text(self):
        from pptx import Presentation

        deck = {"title": "الدوال", "lecturer": "د. سارة",
                "slides": [{"title": "مقدمة", "bullets": ["نقطة أولى"], "notes": ""}]}
        prs = Presentation(str(build_presentation(self.folder / "ar.pptx", deck)))
        self.assertEqual(len(prs.slides), 2)


class MemoryTests(unittest.TestCase):
    """Uses a tiny fake store, so the logic is tested without LangGraph or Postgres."""

    class FakeItem:
        def __init__(self, value):
            self.value = value

    class FakeStore:
        def __init__(self):
            self.data = {}

        def put(self, namespace, key, value):
            self.data[(namespace, key)] = value

        def get(self, namespace, key):
            value = self.data.get((namespace, key))
            return MemoryTests.FakeItem(value) if value else None

        def search(self, namespace, limit=10):
            return [MemoryTests.FakeItem(v) for (ns, _), v in self.data.items() if ns == namespace][:limit]

    def test_saved_session_is_loaded_for_the_same_lecturer_only(self):
        from db.long_term_store import load_memory, save_session

        store = self.FakeStore()
        brief = {"session_id": "s1", "lecturer_id": "Dr. Sara", "topic": "Decorators",
                 "language": "English", "programming_language": "Python", "student_level": "beginner"}
        self.assertTrue(save_session(store, brief, ["Closures", "Syntax"], ["Fewer words"]))

        lines = load_memory(store, "  dr.  sara ")  # spacing and case must not matter
        self.assertTrue(any("English" in line for line in lines))
        self.assertIn("Earlier change request: Fewer words", lines)
        self.assertTrue(any("Decorators" in line for line in lines))
        self.assertEqual(load_memory(store, "Dr. Omar"), [])

    def test_no_store_is_not_an_error(self):
        from db.long_term_store import load_memory, save_session

        self.assertEqual(load_memory(None, "x"), [])
        self.assertFalse(save_session(None, {"lecturer_id": "x", "session_id": "s"}, [], []))


if __name__ == "__main__":
    unittest.main()
